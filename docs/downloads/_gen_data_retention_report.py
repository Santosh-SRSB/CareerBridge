"""CareerBridge data-retention report (DOCX + PDF). Clean copy for sharing."""
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
import win32com.client

NAVY = RGBColor(0x10, 0x13, 0x7C)
OUT_DIRS = [
    Path(r"C:\Users\SRSB HR SOLUTIONS\CareerBridge\docs\downloads"),
    Path(r"C:\Users\SRSB HR SOLUTIONS\Downloads"),
]
NAME = "CareerBridge-Data-Retention-Report"


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def main() -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)
        section.header.is_linked_to_previous = False
        section.header.paragraphs[0].text = ""
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.text = ""
        footer.add_run("Page ")
        add_page_field(footer)

    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    def H(text: str, level: int = 1) -> None:
        heading = doc.add_heading(text, level)
        for run in heading.runs:
            run.font.color.rgb = NAVY

    def P(text: str) -> None:
        doc.add_paragraph(text)

    def B(items: list[str]) -> None:
        for item in items:
            doc.add_paragraph(item, style="List Bullet")

    def T(headers: list[str], rows: list[list[str]]) -> None:
        table = doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.style = "Table Grid"
        for index, header in enumerate(headers):
            cell = table.rows[0].cells[index]
            cell.text = header
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True
                    run.font.size = Pt(9)
                    run.font.color.rgb = NAVY
        for row_index, row in enumerate(rows):
            for col_index, value in enumerate(row):
                cell = table.rows[row_index + 1].cells[col_index]
                cell.text = str(value)
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(9)
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    brand = title.add_run("SRSB CareerBridge")
    brand.bold = True
    brand.font.size = Pt(16)
    brand.font.color.rgb = NAVY

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name = heading.add_run("Data Retention Report")
    name.bold = True
    name.font.size = Pt(18)
    name.font.color.rgb = NAVY

    dated = doc.add_paragraph()
    dated.alignment = WD_ALIGN_PARAGRAPH.CENTER
    dated.add_run(date.today().strftime("%d %B %Y"))

    H("1. Overall data retention chart")
    T(
        ["Category", "Examples", "Kept for", "Removed when"],
        [
            [
                "Candidate account",
                "Login, profile, education, skills, experience, search history, photo",
                "Life of the account",
                "Super Admin permanent delete (Suspended or Inactive only)",
            ],
            [
                "Resume and ATS",
                "Resume record, upload, PDF, ATS report, rewrite, search chunks",
                "365 days from creation",
                "Automatic delete after 365 days (or earlier with account/resume delete)",
            ],
            [
                "Employer and company",
                "Login, company profile, GSTIN / PAN / CIN, logo",
                "Life of the employer account",
                "Super Admin permanent delete (blocked if any payment is above ₹0)",
            ],
            [
                "Jobs and hiring",
                "Jobs, applications, interviews, shortlists, match scores, outcomes",
                "While linked accounts exist",
                "Candidate delete and/or employer delete",
            ],
            [
                "WhatsApp",
                "Messages, phones, template, delivery times",
                "90 days (selection / rejection window)",
                "Automatic delete after 90 days",
            ],
            [
                "AI interactions",
                "Model, tokens, cost, operation log",
                "90 days",
                "Automatic delete after 90 days",
            ],
            [
                "Audit logs",
                "Action, IP, browser, deletion snapshots",
                "365 days",
                "Automatic delete after 365 days",
            ],
            [
                "In-app notifications",
                "Title, body, link, read or unread, device tokens",
                "Until the user account is deleted",
                "Account delete",
            ],
            [
                "Interviews and assessments",
                "AI mock, human mock, skill assessment, purchases, embeddings",
                "Until the candidate account is deleted",
                "Candidate account permanent delete",
            ],
            [
                "Login codes and temporary links",
                "OTP, access and refresh tokens, signed file links, WhatsApp interview link",
                "OTP 5 minutes; access 7 days; refresh 30 days; WhatsApp link 7 days",
                "Code or link expires; underlying records follow their store rules",
            ],
            [
                "Platform configuration",
                "Settings, catalogs, skill dictionary, notification templates, staff logins",
                "Platform life",
                "Admin change or staff delete rules",
            ],
        ],
    )

    H("2. Overall rule")
    P(
        "Candidate and employer account profiles remain for the life of the account. "
        "Timed automatic deletion applies to resumes (365 days), WhatsApp records (90 days), "
        "AI interaction logs (90 days), and audit logs (365 days)."
    )
    P(
        "Permanent account removal occurs only when a Super Admin deletes a Suspended or Inactive account. "
        "An Active account cannot be deleted. An employer with any payment above ₹0 cannot be deleted."
    )

    H("3. Master list — every store")
    T(
        ["Store", "What it holds", "How long it stays"],
        [
            ["Candidate login", "Phone, email, password hash, status, last login", "Until the candidate account is permanently deleted"],
            ["Candidate profile", "Name, location, date of birth, education, skills, experience, salary, notice period, photo, career gaps, WhatsApp opt-in", "Until the candidate account is permanently deleted"],
            ["Job-search history", "Roles and keywords the candidate searched", "Until the candidate account is permanently deleted"],
            ["Resume record", "Title, content, raw text, score, versions", "365 days from creation, then automatic delete"],
            ["Resume files", "Original upload and generated PDF in cloud storage", "Same as the resume record (365 days)"],
            ["ATS report, facts, issues", "Score breakdown, parsed facts, recommendations", "Same as the resume record (365 days)"],
            ["AI resume rewrite", "Before and after score, change list, new version", "Same as the resume record (365 days)"],
            ["Resume search chunks", "Resume text used for matching", "Removed with the resume or candidate account"],
            ["In-app notifications", "Title, body, link, read or unread, for that user", "Until that user's account is permanently deleted"],
            ["Notification templates", "Admin wording for notification text", "Platform data. Kept until an admin deletes the template"],
            ["Device tokens", "Push token for web or mobile", "Until the account is deleted, or the token is replaced"],
            ["WhatsApp messages", "Phones, template, payload, delivery status and times", "90 days, then automatic delete"],
            ["WhatsApp action link", "Secure link inside an interview message", "Link stops working after 7 days; message follows 90-day retention"],
            ["Saved jobs", "Jobs a candidate bookmarked", "Until the candidate or that job is deleted"],
            ["Jobs", "Description, salary, skills, screening questions, status, views", "Until the employer account is permanently deleted. Closing a job does not delete it"],
            ["Applications", "Status, screening answers, private employer note", "Until the candidate or the job is deleted"],
            ["Match scores", "Skills, experience, interview readiness, semantic score", "Removed with the candidate, or with the employer's jobs"],
            ["Employer interviews", "Time, mode, meeting link, notes, feedback, reschedule", "Until the candidate or the employer is deleted"],
            ["Hiring outcomes", "Hired, offer, rejected, withdrawn", "Removed with the employer, unless a charged payment blocks deletion"],
            ["Shortlists", "Employer's private shortlist and note", "Until the candidate, the job, or the employer is deleted"],
            ["Profile view log", "One row per employer, candidate, and month", "Until the candidate or the employer is deleted"],
            ["Employer login and profile", "Company, contact, website, logo, verification", "Until the employer account is permanently deleted"],
            ["GSTIN, PAN, CIN", "Encrypted on the employer row", "Until the employer account is permanently deleted"],
            ["Company GSTIN record", "Encrypted GSTIN and company name, shared by staff", "Can remain after one employer login is deleted"],
            ["GST verification log", "Each verification attempt and result", "Kept. Not removed when the employer is deleted"],
            ["Employer payments of ₹0", "Free payment rows", "Removed with the employer"],
            ["Employer payments above ₹0", "Charged hiring fees", "Kept. They also block deletion of the employer"],
            ["AI mock interview", "Questions, answers, transcript, scores, report", "Until the candidate account is permanently deleted"],
            ["Human mock interview", "Schedule, names, emails, transcript, fee, score", "Until the candidate account is permanently deleted"],
            ["Skill assessment", "Questions, answers, score, feedback", "Until the candidate account is permanently deleted"],
            ["Assessment purchases", "Credits bought and amount paid by the candidate", "Until the candidate account is permanently deleted"],
            ["AI cost log", "Model, token counts, cost, error. Resume text is not stored here", "90 days, then automatic delete"],
            ["Testimonials", "Rating, quote, name, approval status, prompts", "Until that user's account is permanently deleted"],
            ["OTP requests", "Hashed login code, phone or email, attempt count", "The code works for 5 minutes. The row stays"],
            ["Refresh tokens", "Hashed session token", "Works for 30 days by default. The row stays until the account is deleted"],
            ["Audit log", "Action, IP address, browser, and on deletion the name and email", "365 days, then automatic delete"],
            ["Platform settings and catalogs", "Admin configuration, cities, categories, skill names", "Platform data. Not a person's record"],
            ["Cloud SQL backups", "A copy of the database at backup time", "Set on the production Cloud SQL instance"],
        ],
    )

    H("4. Candidate account")
    P(
        "The following data remains for the life of the candidate account "
        "(Active, Suspended, or Inactive). Suspension hides access; it does not erase data."
    )
    T(
        ["Data", "Contents"],
        [
            ["Login", "Phone, email, password hash, account status, last login time"],
            ["Profile", "Name, city, state, preferred work city, date of birth, gender, about, photo, profile links"],
            ["Career", "Education, skills, experience, career gaps, certifications, projects"],
            ["Preferences", "Notice period, employment status, expected salary, job types, relocation, language"],
            ["WhatsApp on the profile", "Opt-in, opt-in time, number, verified flag"],
            ["Search history", "Each role or keyword searched, and when it was last searched"],
            ["Onboarding", "Completed flag, skipped steps, profile completion, dashboard reached"],
        ],
    )
    P(
        "When Super Admin permanently deletes a Suspended or Inactive candidate, the login and "
        "related records are removed together: profile, education, skills, experience, career gaps, "
        "search history, applications, interviews, resumes, assessments, saved jobs, match scores, "
        "and search embeddings. Profile photo and resume files are then removed from cloud storage."
    )

    H("5. Resume data")
    P("A resume is stored in two places: the database row, and the file in private cloud storage.")
    T(
        ["Piece", "What is kept", "When it is removed"],
        [
            ["Resume row", "Title, target job, template, summary, full structured content, raw extracted text, score, version, kind", "Automatic delete 365 days after creation; also on candidate account delete"],
            ["Original file", "The file the candidate uploaded", "Same as the resume row"],
            ["Generated PDF", "The PDF CareerBridge builds from the resume", "Same as the resume row"],
            ["ATS report", "Overall score, section scores, issue list, recommended plan", "With the resume"],
            ["Facts and issues", "Each parsed fact and each ATS issue, with the recommendation", "With the resume"],
            ["Rewrite history", "Plan, before score, after score, improvements, the resulting version", "With the source resume"],
            ["Search chunks", "Sections of the resume text and their vectors, used for job matching", "With the resume or candidate account"],
        ],
    )
    P(
        "If the candidate deletes a resume from the product before day 365, unused resumes are removed at once "
        "(including cloud files). Resumes linked to applications are archived until the 365-day automatic delete "
        "or candidate account deletion."
    )

    H("6. Employer and company data")
    T(
        ["Data", "How long"],
        [
            ["Login: phone, email, password hash, status", "Until the employer account is permanently deleted"],
            ["Profile: company name, industry, city, state, contact, designation, website, about, LinkedIn, size, logo, verification status", "Until the employer account is permanently deleted"],
            ["GSTIN, PAN, and CIN", "Encrypted on the employer row, until that account is deleted"],
            ["Shared company record", "Encrypted GSTIN and legal name. Several staff accounts can point at one company. Deleting one employer login leaves this company row in place"],
            ["GST verification attempts", "Each attempt is kept: encrypted GSTIN, pass or fail, provider, and time. This log is not removed with the employer"],
            ["Logo file", "Removed when the employer account is permanently deleted"],
            ["Payments of ₹0", "Removed with the employer"],
            ["Payments above ₹0", "Kept. Super Admin cannot delete an employer who has any charged payment, so the jobs and profile stay with the payment history"],
        ],
    )

    H("7. Jobs, applications, and hiring")
    P(
        "Closing, pausing, or filling a job records a status and a closed time. The job row stays. "
        "There is one application per candidate per job."
    )
    T(
        ["Record", "Kept while", "Removed when"],
        [
            ["Job", "The employer account exists", "The employer account is permanently deleted"],
            ["Application and screening answers", "Both the candidate and the job exist", "The candidate is deleted, or the employer's jobs are deleted"],
            ["Private employer note on an application", "The application exists", "The application is removed"],
            ["Match score and embeddings for that job", "The candidate and the job exist", "The candidate is deleted, or the employer's jobs are deleted"],
            ["Employer interview", "The application exists", "The candidate or the employer is deleted"],
            ["Hiring outcome", "The employer account can still be deleted", "The employer is deleted. A charged payment blocks that delete"],
            ["Talent shortlist", "Candidate, job, and employer exist", "Any one of those three is deleted"],
            ["Monthly profile-view count", "Candidate and employer exist", "The candidate or the employer is deleted"],
            ["Saved job", "Candidate and job exist", "The candidate or the job is deleted"],
        ],
    )
    P(
        "Deleting a candidate removes that person's applications. The employer's jobs stay, and the "
        "employer's payment records stay. Deleting an employer removes that employer's jobs and the "
        "applications to those jobs. The candidate accounts themselves stay."
    )

    H("8. WhatsApp")
    P(
        "WhatsApp records are kept for 90 days to cover employer selection or rejection confirmation. "
        "Messages older than 90 days are automatically deleted."
    )
    T(
        ["What is stored", "Detail"],
        [
            ["Phones", "Destination number and, for inbound messages, the sender number"],
            ["Message", "Type, template name, direction (outbound or inbound), and the payload"],
            ["Delivery", "Status: queued, sent, delivered, read, failed, or received"],
            ["Times", "Created, sent, delivered, and read"],
            ["Provider id", "WhatsApp message id, plus any error payload"],
            ["Links", "Optional link to a candidate and to an employer interview"],
            ["Retention", "90 days from message creation, then automatic delete"],
        ],
    )
    P(
        "The secure link inside an interview WhatsApp message expires after 7 days. "
        "The message itself follows the 90-day retention rule."
    )

    H("9. In-app notifications")
    P(
        "Each notification is stored for one user (title, body, type, link, read time). "
        "Marking as read does not delete it. There is no cleanup of old or read notifications."
    )
    T(
        ["Store", "How long"],
        [
            ["Notification sent to a candidate or employer", "Until that user's account is permanently deleted. Read and unread rows are both kept"],
            ["Push device token", "Until the account is deleted, or until logout replaces that token"],
            ["Notification text templates", "Platform wording managed by admins. Kept until an admin deletes that template. Not tied to a user"],
        ],
    )
    P(
        "OTP and interview emails are sent through the mail server. "
        "CareerBridge does not store a mailbox copy; the OTP hash, notification, or interview record is kept instead."
    )

    H("10. Interviews, assessments, and AI")
    T(
        ["Store", "Contents", "How long"],
        [
            ["AI mock interview", "Role, questions, answers, transcript, warnings, scores, written report", "Until the candidate account is permanently deleted"],
            ["Human mock interview", "Schedule, join token, interviewer and candidate name and email, transcript, fee, score, feedback", "Until the candidate account is permanently deleted"],
            ["Skill assessment", "Skills, questions, answers, score, feedback", "Until the candidate account is permanently deleted"],
            ["Assessment credit purchase", "Credits, amount in rupees, payment status", "Until the candidate account is permanently deleted"],
            ["AI cost log", "Operation name, model, input and output tokens, latency, cost, error text", "90 days, then automatic delete. Resume or answer text is not stored here"],
            ["Profile and resume embeddings", "Vectors used to match candidates and jobs", "Removed when that candidate, resume, or employer's jobs are deleted"],
        ],
    )

    H("11. Login codes, sessions, and temporary links")
    P("These limits expire access only. They do not delete candidate, employer, resume, or WhatsApp records.")
    T(
        ["Item", "Works for", "What remains afterward"],
        [
            ["Email or mobile OTP", "5 minutes, or 5 wrong attempts", "The code stops working. The database keeps the hashed code, phone or email, and attempt count"],
            ["Sign-in access token", "7 days, unless the live service sets JWT_ACCESS_EXPIRES to another value", "The user must sign in again or use the refresh token"],
            ["Refresh token", "30 days by the same default (JWT_REFRESH_EXPIRES)", "The session ends. The token row stays until the account is deleted"],
            ["Registration draft", "30 minutes, in that browser only", "Dropped from the browser. It is never written to the database"],
            ["Resume view link", "30 minutes", "The resume file stays. The next open creates a new link"],
            ["Profile photo link", "7 days", "The photo file stays. The next open creates a new link"],
            ["Company logo link", "60 minutes", "The logo file stays. The next open creates a new link"],
            ["WhatsApp interview link", "7 days", "The interview and the WhatsApp message stay. The link stops opening"],
        ],
    )

    H("12. What stays after an account is deleted")
    T(
        ["Record", "What is still in the system"],
        [
            ["Audit log", "Deletion snapshot (name, email, role, status) remains until the 365-day audit retention ends"],
            ["WhatsApp log", "May remain until the 90-day WhatsApp retention ends; person link is cleared on account delete"],
            ["AI cost log", "May remain until the 90-day AI interaction retention ends"],
            ["GST verification log", "Encrypted GSTIN and the verification result. Not removed by account deletion"],
            ["Charged employer payments", "Amount, status, and provider reference. Their existence blocks employer deletion"],
            ["Shared company GSTIN row", "Can remain when one staff login is removed and other staff still use the company"],
            ["Cloud SQL backup", "A deleted row can still exist inside a backup taken before the delete, until that backup ages out"],
        ],
    )
    P(
        "After a candidate applies, the employer may also keep a copy outside CareerBridge. "
        "Retention of that copy is outside CareerBridge control."
    )

    H("13. Platform data")
    P("These are configuration records, not candidate or employer personal files.")
    B(
        [
            "Platform settings: admin keys and values.",
            "Catalogs: job categories, cities, experience levels, and languages.",
            "Skill dictionary: shared skill names and categories.",
            "Notification templates: default and overridden notification wording.",
            "Staff accounts: admin and operator logins. A Super Admin account cannot be deleted. An admin or operator can be deleted only after that staff account is suspended or inactive. Actions they performed stay in the audit log.",
        ]
    )

    H("14. How permanent deletion works")
    B(
        [
            "Only Super Admin can permanently delete an account.",
            "The account must already be Suspended or Inactive. An Active account is refused.",
            "Candidate delete removes the person, their resumes and files, applications, interviews, assessments, notifications, sessions, and match data.",
            "Employer delete removes the company login, jobs, interviews, and applications to those jobs. Candidate accounts stay.",
            "Employer delete is refused when any payment amount is above ₹0.",
            "A staff member cannot delete their own account. Super Admin accounts cannot be deleted.",
            "The audit log of the deletion is written in the same transaction and is kept.",
        ]
    )

    H("15. Backups")
    P(
        "Automated Cloud SQL backups retain a copy of the database for the configured backup window. "
        "Until a backup expires, it may still contain rows that were later deleted from the live database."
    )

    H("16. Summary")
    B(
        [
            "Resumes: kept 365 days, then automatically deleted.",
            "WhatsApp records: kept 90 days (selection / rejection window), then automatically deleted.",
            "AI interactions: kept 90 days, then automatically deleted.",
            "Audit logs: kept 365 days, then automatically deleted.",
            "Candidate and employer account profiles remain until Super Admin permanent delete.",
            "Suspended and inactive accounts retain profile data until permanent delete.",
            "An employer with any charged payment cannot be deleted.",
            "Cloud SQL backups can outlive a delete for the configured backup window.",
        ]
    )

    for folder in OUT_DIRS:
        folder.mkdir(parents=True, exist_ok=True)
        docx_path = folder / f"{NAME}.docx"
        doc.save(docx_path)
        print("DOCX", docx_path, docx_path.stat().st_size)

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        for folder in OUT_DIRS:
            docx_path = folder / f"{NAME}.docx"
            pdf_path = folder / f"{NAME}.pdf"
            opened = word.Documents.Open(str(docx_path))
            opened.SaveAs(str(pdf_path), FileFormat=17)
            opened.Close(False)
            print("PDF", pdf_path, pdf_path.stat().st_size)
    finally:
        word.Quit()


if __name__ == "__main__":
    main()
